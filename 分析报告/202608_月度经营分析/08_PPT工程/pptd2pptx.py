# -*- coding: utf-8 -*-
"""PPTD -> PPTX 本地机械转译器（仅使用特性子集）。"""
import os
import re
import sys
import yaml
from html.parser import HTMLParser
from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

EMU_PER_PT = 12700
PAGE_W, PAGE_H = 960, 540

ALIGN_H = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}
ANCHOR_V = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}


def to_pt(v):
    return float(v)


def resolve_color(value, theme):
    if value is None:
        return '#000000'
    if isinstance(value, str) and value.startswith('$'):
        key = value[1:]
        return theme.get('colors', {}).get(key, '#000000')
    if isinstance(value, str) and not value.startswith('#'):
        return '#' + value
    return value


def rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip('#'))


def load_deck(pptd_dir):
    path = os.path.join(pptd_dir, 'deck.pptd')
    with open(path, 'r', encoding='utf-8') as f:
        return yaml.safe_load(f)


def load_page(path):
    with open(path, 'r', encoding='utf-8') as f:
        return yaml.safe_load(f)


def blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name == 'Blank':
            return layout
    return prs.slide_layouts[6]


def set_run_font(run, name, size_pt=None, bold=None, italic=None, color=None):
    font = run.font
    if name:
        font.name = name
        rpr = run._r.get_or_add_rPr()
        ea = rpr.find(qn('a:ea'))
        if ea is None:
            ea = OxmlElement('a:ea')
            rpr.append(ea)
        ea.set('typeface', name)
    if size_pt is not None:
        font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if color is not None:
        font.color.rgb = rgb(color)


class RichTextParser(HTMLParser):
    def __init__(self, default):
        super().__init__()
        self.default = default
        self.paragraphs = []
        self.stack = [{}]
        self._para = None

    def _style(self):
        merged = dict(self.default)
        for s in self.stack[1:]:
            for k, v in s.items():
                if v is not None:
                    merged[k] = v
        return merged

    def _ensure_para(self, pstyle=None):
        if self._para is None:
            self._para = {'style': pstyle or {}, 'runs': []}
            self.paragraphs.append(self._para)

    def _parse_style(self, raw):
        out = {}
        for part in raw.split(';'):
            if ':' not in part:
                continue
            k, v = part.split(':', 1)
            k, v = k.strip(), v.strip()
            if k == 'text-align':
                out['align'] = v
            elif k == 'line-height':
                out['lineHeight'] = float(v)
            elif k == 'margin-top':
                out['marginTop'] = float(v.replace('px', ''))
            elif k == 'color':
                out['color'] = v
            elif k == 'font-size':
                out['fontSize'] = float(v.replace('px', ''))
        return out

    def handle_starttag(self, tag, attrs):
        ad = dict(attrs)
        if tag == 'p':
            self._para = None
            self._ensure_para(self._parse_style(ad.get('style', '')))
        elif tag == 'span':
            self.stack.append(self._parse_style(ad.get('style', '')))
        elif tag == 'strong':
            self.stack.append({'bold': True})
        elif tag == 'em':
            self.stack.append({'italic': True})
        elif tag == 'br':
            self._ensure_para()
            assert self._para is not None
            self._para['runs'].append({**self._style(), 'text': '\n'})

    def handle_endtag(self, tag):
        if tag in ('span', 'strong', 'em') and len(self.stack) > 1:
            self.stack.pop()

    def handle_data(self, data):
        if data:
            self._ensure_para()
            assert self._para is not None
            self._para['runs'].append({**self._style(), 'text': data})


def make_plain_paragraphs(text, default):
    paras = []
    for line in text.splitlines():
        paras.append({'style': {}, 'runs': [{**default, 'text': line}]})
    return paras


def apply_paragraph(p, para, default, theme):
    halign = para['style'].get('align') or (default.get('align') or ['left', 'top'])[0]
    p.alignment = ALIGN_H.get(halign, PP_ALIGN.LEFT)
    lh = para['style'].get('lineHeight') or default.get('lineHeight')
    if lh:
        p.line_spacing = float(lh)
    mt = para['style'].get('marginTop', 0)
    if mt:
        p.space_before = Pt(mt)


def apply_run(run, info, default, theme):
    color = info.get('color') or default.get('color')
    size = info.get('fontSize') if info.get('fontSize') is not None else default.get('fontSize')
    bold = info.get('bold') if info.get('bold') is not None else default.get('bold')
    italic = info.get('italic') if info.get('italic') is not None else default.get('italic')
    name = default.get('fontFamily', 'HarmonyOS Sans SC')
    set_run_font(run, name, size, bold, italic, resolve_color(color, theme))


def add_text(slide, elem, theme):
    x, y, w, h = elem['bounds']
    default = elem.get('content', {})
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    valign = (default.get('align') or ['left', 'top'])[1]
    tf.vertical_anchor = ANCHOR_V.get(valign, MSO_ANCHOR.TOP)
    text = default.get('text', '')
    if '<' in text:
        parser = RichTextParser(default)
        parser.feed(text)
        paras = parser.paragraphs or [{'style': {}, 'runs': []}]
    else:
        paras = make_plain_paragraphs(text, default)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.clear()
        apply_paragraph(p, para, default, theme)
        for info in para['runs']:
            run = p.add_run()
            run.text = info['text']
            apply_run(run, info, default, theme)
    return shape


def set_shape_gradient(shape, angle, stops, theme):
    spPr = shape._sp.spPr
    for tag in ('noFill', 'solidFill', 'gradFill', 'blipFill', 'pattFill', 'grpFill'):
        el = spPr.find(qn('a:' + tag))
        if el is not None:
            spPr.remove(el)
    ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    stop_xml = ''.join(
        '<a:gs pos="%d"><a:srgbClr val="%s"/></a:gs>' % (
            int(s['position'] * 100000),
            resolve_color(s['color'], theme).lstrip('#')
        ) for s in stops
    )
    xml = '<a:gradFill xmlns:a="%s" rotWithShape="1"><a:gsLst>%s</a:gsLst><a:lin ang="%d" scaled="1"/></a:gradFill>' % (
        ns, stop_xml, int(angle * 60000))
    spPr.append(parse_xml(xml))


def add_shape(slide, elem, theme):
    x, y, w, h = elem['bounds']
    name = elem.get('shapeName', 'rect')
    mso = MSO_SHAPE.ROUNDED_RECTANGLE if name == 'roundRect' else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(mso, Pt(x), Pt(y), Pt(w), Pt(h))
    if name == 'roundRect' and elem.get('adjustments'):
        shape.adjustments[0] = elem['adjustments'][0] / 100000.0
    fill = elem.get('fill')
    if fill and fill.get('type') == 'solid':
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(resolve_color(fill['color'], theme))
    elif fill and fill.get('type') == 'gradient':
        set_shape_gradient(shape, fill.get('angle', 0), fill.get('stops', []), theme)
    else:
        shape.fill.background()
    border = elem.get('border')
    if border:
        shape.line.color.rgb = rgb(resolve_color(border['color'], theme))
        shape.line.width = Pt(border['width'])
    else:
        shape.line.fill.background()
    return shape


def image_size(src):
    with PILImage.open(src) as im:
        return im.size


def add_image(slide, elem, pptd_dir):
    x, y, w, h = elem['bounds']
    src = os.path.join(pptd_dir, elem['src'])
    iw, ih = image_size(src)
    mode = elem.get('fit', {}).get('mode', 'fill')
    if mode == 'fill':
        return slide.shapes.add_picture(src, Pt(x), Pt(y), Pt(w), Pt(h))
    if mode == 'contain':
        scale = min(w / iw, h / ih)
    else:
        scale = max(w / iw, h / ih)
    sw, sh = iw * scale, ih * scale
    px, py = x + (w - sw) / 2.0, y + (h - sh) / 2.0
    pic = slide.shapes.add_picture(src, Pt(px), Pt(py), width=Pt(sw))
    if mode == 'cover':
        pic.crop_left = max(0.0, (x - px) / sw)
        pic.crop_top = max(0.0, (y - py) / sh)
        pic.crop_right = max(0.0, (px + sw - (x + w)) / sw)
        pic.crop_bottom = max(0.0, (py + sh - (y + h)) / sh)
    return pic


def set_cell_border(cell, color_hex, width_pt):
    tcPr = cell._tc.get_or_add_tcPr()
    wEmu = int(width_pt * EMU_PER_PT)
    cval = color_hex.lstrip('#')
    for tag in ('lnL', 'lnR', 'lnT', 'lnB'):
        ln = tcPr.find(qn('a:' + tag))
        if ln is None:
            ln = OxmlElement('a:' + tag)
            tcPr.append(ln)
        ln.set('w', str(wEmu))
        ln.set('cap', 'flat')
        fill = ln.find(qn('a:solidFill'))
        if fill is None:
            fill = OxmlElement('a:solidFill')
            ln.append(fill)
        for child in list(fill):
            fill.remove(child)
        srgb = OxmlElement('a:srgbClr')
        srgb.set('val', cval)
        fill.append(srgb)


def set_cell_text(cell, text, defaults, theme):
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.clear()
    p.alignment = ALIGN_H.get(defaults.get('align', ['center', 'middle'])[0], PP_ALIGN.CENTER)
    if '<' in text:
        parser = RichTextParser(defaults)
        parser.feed(text)
        para = parser.paragraphs[0] if parser.paragraphs else {'style': {}, 'runs': []}
        apply_paragraph(p, para, defaults, theme)
        for info in para['runs']:
            run = p.add_run()
            run.text = info['text']
            apply_run(run, info, defaults, theme)
    else:
        run = p.add_run()
        run.text = text
        apply_run(run, defaults, defaults, theme)


def add_table(slide, elem, theme):
    x, y, w, h = elem['bounds']
    rows = elem['rows']
    n_rows = len(rows)
    n_cols = max(len(r) for r in rows)
    shape = slide.shapes.add_table(n_rows, n_cols, Pt(x), Pt(y), Pt(w), Pt(h))
    tbl = shape.table
    tbl.first_row = tbl.last_row = tbl.first_col = tbl.last_col = False
    tbl.band_rows = tbl.band_cols = tbl.horz_banding = tbl.vert_banding = False
    style = theme['tableStyles']['default']
    cws = elem['columnWidths']
    rhs = elem['rowHeights']
    csum = sum(cws)
    rsum = sum(rhs)
    for i, p in enumerate(cws):
        tbl.columns[i].width = int(w * EMU_PER_PT * p / csum)
    for i, p in enumerate(rhs):
        tbl.rows[i].height = int(h * EMU_PER_PT * p / rsum)
    cell_defaults = style['cellStyle']
    border = cell_defaults.get('border', {})
    bcolor = resolve_color(border.get('color', '$border'), theme)
    bwidth = border.get('width', 1)
    for r_idx, row in enumerate(rows):
        for c_idx, cell_data in enumerate(row):
            _style_table_cell(tbl.cell(r_idx, c_idx), r_idx, cell_data,
                              style, theme, cell_defaults, bcolor, bwidth)
    return shape


def _style_table_cell(cell, r_idx, cell_data, style, theme, cell_defaults, bcolor, bwidth):
    cell.fill.solid()
    if r_idx == 0:
        fill_color = style['firstRowStyle']['fill']['color']
        font_color = style['firstRowStyle'].get('color', '#000000')
        bold = style['firstRowStyle'].get('bold', True)
        size = style['firstRowStyle'].get('fontSize', cell_defaults.get('fontSize'))
    else:
        bi = (r_idx - 1) % len(style['bodyStyles'])
        fill_color = style['bodyStyles'][bi]['fill']['color']
        font_color = cell_data.get('color') or cell_defaults.get('color', '$text')
        bold = cell_data.get('bold')
        size = cell_data.get('fontSize') or cell_defaults.get('fontSize')
    cell.fill.fore_color.rgb = rgb(resolve_color(fill_color, theme))
    valign = (cell_data.get('align') or cell_defaults.get('align') or ['center', 'middle'])[1]
    cell.vertical_anchor = ANCHOR_V.get(valign, MSO_ANCHOR.MIDDLE)
    defaults = {
        'fontFamily': cell_defaults.get('fontFamily', 'HarmonyOS Sans SC'),
        'fontSize': size,
        'bold': bold,
        'italic': False,
        'color': resolve_color(font_color, theme),
        'align': cell_data.get('align') or cell_defaults.get('align') or ['center', 'middle'],
    }
    set_cell_text(cell, str(cell_data.get('text', '')), defaults, theme)
    set_cell_border(cell, bcolor, bwidth)


def add_element(slide, elem, theme, pptd_dir):
    etype = elem.get('elementType')
    if etype == 'text':
        return add_text(slide, elem, theme)
    if etype == 'shape':
        return add_shape(slide, elem, theme)
    if etype == 'image':
        return add_image(slide, elem, pptd_dir)
    if etype == 'table':
        return add_table(slide, elem, theme)
    raise ValueError('不支持的元素类型: %s' % etype)


def add_background_image(slide, bg, pptd_dir):
    src = os.path.join(pptd_dir, bg['src'])
    return slide.shapes.add_picture(src, 0, 0, Pt(PAGE_W), Pt(PAGE_H))


def apply_background(slide, bg, pptd_dir, theme):
    if bg and bg.get('type') == 'solid':
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(resolve_color(bg['color'], theme))
    else:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb('#FFFFFF')


def main():
    base = os.path.dirname(os.path.abspath(__file__))
    pptd_dir = os.path.join(base, 'pptd')
    out_path = os.path.join(base, '2026年8月经营分析汇报.pptx')
    deck = load_deck(pptd_dir)
    theme = deck.get('theme', {})
    prs = Presentation()
    prs.slide_width = Emu(PAGE_W * EMU_PER_PT)
    prs.slide_height = Emu(PAGE_H * EMU_PER_PT)
    layout = blank_layout(prs)
    for page_path in deck.get('pages', []):
        page = load_page(os.path.join(pptd_dir, page_path))
        slide = prs.slides.add_slide(layout)
        bg = page.get('background')
        apply_background(slide, bg, pptd_dir, theme)
        if bg and bg.get('type') == 'image':
            add_background_image(slide, bg, pptd_dir)
        for elem in page.get('elements', []):
            add_element(slide, elem, theme, pptd_dir)
    prs.save(out_path)
    print('已生成: %s' % out_path)


if __name__ == '__main__':
    main()
