"""提取 pptx 每页元素几何与字号，供版式规范对标。"""
import json
import sys

from pptx import Presentation
from pptx.util import Emu


def dump(path: str) -> list:
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        shapes = []
        for sh in slide.shapes:
            item = {
                "type": str(sh.shape_type),
                "name": sh.name,
                "l": round(Emu(sh.left).inches, 2) if sh.left is not None else None,
                "t": round(Emu(sh.top).inches, 2) if sh.top is not None else None,
                "w": round(Emu(sh.width).inches, 2) if sh.width is not None else None,
                "h": round(Emu(sh.height).inches, 2) if sh.height is not None else None,
            }
            if sh.has_text_frame:
                sizes = []
                txt = ""
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            sizes.append(round(r.font.size.pt, 1))
                        txt += r.text
                item["font_sizes"] = sorted(set(sizes))
                item["text"] = txt[:40].replace("\n", " ")
            shapes.append(item)
        out.append({"slide": i, "shapes": shapes})
    return out


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(dump(src), f, ensure_ascii=False, indent=1)
    print(f"OK {dst}")
