# -*- coding: utf-8 -*-
"""平台 Excel 侦察：加密头/sheet 名/表头列名（对比 quarterly 脚本字段映射需求）"""
import sys, os
sys.stdout.reconfigure(encoding='utf-8')

path = r'E:\3-其他资料\数据分析\sales_analytics_platform\data\出货明细-8月（9.5）.xlsx'
head = open(path, 'rb').read(4)
print('文件头:', head, '| PK头(明文):', head[:2] == b'PK')

if head[:2] == b'PK':
    try:
        from python_calamine import CalamineWorkbook
        wb = CalamineWorkbook.from_path(path)
        names = wb.sheet_names
        print('sheet 列表:', names)
        sh = wb.get_sheet_by_name(names[0])
        rows = sh.to_python(skip_empty_area=False)
        print('首 sheet 行数:', len(rows))
        hdr = rows[0] if rows else []
        print('表头列数:', len(hdr))
        print('表头:', [str(c)[:18] for c in hdr])
    except Exception as e:
        print('calamine 读取失败:', type(e).__name__, str(e)[:200])
        try:
            import openpyxl
            wb2 = openpyxl.load_workbook(path, read_only=True)
            print('openpyxl sheet 列表:', wb2.sheetnames)
            ws = wb2[wb2.sheetnames[0]]
            row1 = next(ws.iter_rows(max_row=1, values_only=True))
            print('表头列数:', len(row1), '| 前30列:', [str(c)[:18] for c in row1[:30]])
            wb2.close()
        except Exception as e2:
            print('openpyxl 也失败:', type(e2).__name__, str(e2)[:200])
else:
    print('非 PK 头 = DSE 加密，Python 不可直读；需 COM 转存明文副本')
