# -*- coding: utf-8 -*-
r"""提取7月报告docx纲要：标题层级+段落摘要+表格结构"""
import docx

DOC = r"\\192.168.8.3\财务部\财务电子档案备份\D1经营分析\分析报告\分析报告-202608\2026年7月销售经营分析报告(5).docx"
OUT = r"C:\Users\910373\AppData\Local\Temp\opencode\build\report_outline.txt"

d = docx.Document(DOC)
lines = []
lines.append(f"段落数: {len(d.paragraphs)}  表格数: {len(d.tables)}")

def para_text(p):
    return "".join(r.text for r in p.runs).strip() or p.text.strip()

lines.append("\n===== 标题结构(Heading/标题样式) =====")
for i, p in enumerate(d.paragraphs):
    st = p.style.name if p.style else ""
    t = para_text(p)
    if not t:
        continue
    if ("Heading" in st or "标题" in st) and len(t) < 80:
        lines.append(f"[{i}] ({st}) {t}")

lines.append("\n===== 正文段落摘要(非空,截120字) =====")
for i, p in enumerate(d.paragraphs):
    t = para_text(p)
    if t and len(t) > 4:
        lines.append(f"[{i}] {t[:120]}")

lines.append("\n===== 表格结构 =====")
for ti, tb in enumerate(d.tables):
    try:
        nrow = len(tb.rows)
        ncol = len(tb.columns)
        hdr = " | ".join(c.text.strip()[:14] for c in tb.rows[0].cells[:8])
        sample = " | ".join(c.text.strip()[:10] for c in tb.rows[1].cells[:8]) if nrow > 1 else ""
        lines.append(f"表{ti+1}: {nrow}x{ncol} | 表头: {hdr} | 次行: {sample}")
    except Exception as e:
        lines.append(f"表{ti+1}: <读取失败 {e}>")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("paras:", len(d.paragraphs), "tables:", len(d.tables))
print("OUTLINE_DONE")
