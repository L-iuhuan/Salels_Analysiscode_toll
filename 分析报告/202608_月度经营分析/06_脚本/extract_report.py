# -*- coding: utf-8 -*-
import sys, zipfile, re
sys.stdout.reconfigure(encoding='utf-8')
path = r'E:\3-其他资料\数据分析\2026年8月销售经营分析报告(定稿).docx'
lines = []
try:
    import docx
    d = docx.Document(path)
    for p in d.paragraphs:
        if p.text.strip(): lines.append(p.text.strip())
    for t in d.tables:
        lines.append('── 表格 ──')
        for row in t.rows:
            lines.append(' | '.join(c.text.strip() for c in row.cells))
except ImportError:
    z = zipfile.ZipFile(path)
    xml = z.read('word/document.xml').decode('utf-8')
    xml = re.sub(r'</w:p>', '\n', xml)
    text = re.sub(r'<[^>]+>', '', xml)
    import html as H
    text = H.unescape(text)
    lines = [l.strip() for l in text.splitlines() if l.strip()]
out = r'C:\Users\910373\AppData\Local\Temp\opencode\report_202608.txt'
open(out, 'w', encoding='utf-8').write('\n'.join(lines))
print('提取完成: 行数=%d 字符=%d → %s' % (len(lines), sum(len(l) for l in lines), out))
